"""
AI-powered features routes with input validation.

Security features:
- JWT authentication required for all endpoints
- Input length limits to prevent abuse
- Rate limiting via main.py middleware
- User-scoped data access
"""

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from pydantic import BaseModel, Field, field_validator
from typing import Optional
import re
import io

from database import get_db
from models import Job, Company, User
from services import ai_service
from middleware.auth import get_current_user

router = APIRouter(prefix="/api/ai", tags=["ai"])

# Maximum text lengths for AI inputs (to control costs and prevent abuse)
MAX_RESUME_LENGTH = 50000  # ~10 pages
MAX_JOB_DESCRIPTION_LENGTH = 50000
MAX_TITLE_LENGTH = 500

# Valid interview types
VALID_INTERVIEW_TYPES = {"phone_screen", "technical", "behavioral", "onsite", "final"}


class CoverLetterRequest(BaseModel):
    job_id: Optional[int] = Field(None, ge=1)
    job_title: Optional[str] = Field(None, max_length=MAX_TITLE_LENGTH)
    company_name: Optional[str] = Field(None, max_length=255)
    job_description: Optional[str] = Field(None, max_length=MAX_JOB_DESCRIPTION_LENGTH)
    resume_text: str = Field(..., min_length=100, max_length=MAX_RESUME_LENGTH)


class InterviewQuestionsRequest(BaseModel):
    job_id: Optional[int] = Field(None, ge=1)
    job_title: Optional[str] = Field(None, max_length=MAX_TITLE_LENGTH)
    job_description: Optional[str] = Field(None, max_length=MAX_JOB_DESCRIPTION_LENGTH)
    interview_type: str = Field("behavioral", max_length=50)

    @field_validator('interview_type')
    @classmethod
    def validate_interview_type(cls, v):
        if v not in VALID_INTERVIEW_TYPES:
            raise ValueError(f"Invalid interview type. Must be one of: {', '.join(VALID_INTERVIEW_TYPES)}")
        return v


class ResumeMatchRequest(BaseModel):
    job_id: Optional[int] = Field(None, ge=1)
    job_description: Optional[str] = Field(None, max_length=MAX_JOB_DESCRIPTION_LENGTH)
    resume_text: str = Field(..., min_length=100, max_length=MAX_RESUME_LENGTH)


class ResumeImprovementRequest(BaseModel):
    resume_text: str = Field(..., min_length=100, max_length=MAX_RESUME_LENGTH)
    job_id: Optional[int] = Field(None, ge=1)
    job_description: Optional[str] = Field(None, max_length=MAX_JOB_DESCRIPTION_LENGTH)


class CompanyResearchRequest(BaseModel):
    company_id: Optional[int] = Field(None, ge=1)
    company_name: Optional[str] = Field(None, max_length=255)
    industry: Optional[str] = Field(None, max_length=100)
    website: Optional[str] = Field(None, max_length=500)

    @field_validator('website')
    @classmethod
    def validate_website(cls, v):
        if v and not re.match(r'^https?://', v):
            raise ValueError("Website must start with http:// or https://")
        return v


class ATSResumeRequest(BaseModel):
    job_id: Optional[int] = Field(None, ge=1)
    job_title: Optional[str] = Field(None, max_length=MAX_TITLE_LENGTH)
    company_name: Optional[str] = Field(None, max_length=255)
    job_description: Optional[str] = Field(None, max_length=MAX_JOB_DESCRIPTION_LENGTH)
    resume_text: str = Field(..., min_length=100, max_length=MAX_RESUME_LENGTH)


@router.post("/cover-letter")
def generate_cover_letter(
    request: CoverLetterRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Generate a tailored cover letter for a job application."""
    job_title = request.job_title
    company_name = request.company_name
    job_description = request.job_description

    if request.job_id:
        # Allow both user-owned jobs and public jobs (user_id is null)
        job = db.query(Job).filter(Job.id == request.job_id).first()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        # Only check ownership if the job belongs to a user
        if job.user_id is not None and job.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this job")
        job_title = job_title or job.title
        job_description = job_description or job.job_description
        if job.company:
            company_name = company_name or job.company.name

    if not job_title:
        raise HTTPException(status_code=400, detail="Job title is required")
    if not job_description:
        raise HTTPException(status_code=400, detail="Job description is required")
    if not request.resume_text:
        raise HTTPException(status_code=400, detail="Resume text is required")

    try:
        cover_letter = ai_service.generate_cover_letter(
            job_title=job_title,
            company_name=company_name or "the company",
            job_description=job_description,
            resume_text=request.resume_text
        )
        return {"cover_letter": cover_letter}
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate cover letter: {str(e)}")


@router.post("/interview-questions")
def generate_interview_questions(
    request: InterviewQuestionsRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Generate practice interview questions for a job."""
    job_title = request.job_title
    job_description = request.job_description

    if request.job_id:
        # Allow both user-owned jobs and public jobs (user_id is null)
        job = db.query(Job).filter(Job.id == request.job_id).first()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        if job.user_id is not None and job.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this job")
        job_title = job_title or job.title
        job_description = job_description or job.job_description

    if not job_title:
        raise HTTPException(status_code=400, detail="Job title is required")
    if not job_description:
        job_description = f"Position: {job_title}"

    valid_types = ["phone_screen", "technical", "behavioral", "onsite", "final"]
    if request.interview_type not in valid_types:
        raise HTTPException(status_code=400, detail=f"Interview type must be one of: {valid_types}")

    try:
        questions = ai_service.generate_interview_questions(
            job_title=job_title,
            job_description=job_description,
            interview_type=request.interview_type
        )
        return {"questions": questions, "interview_type": request.interview_type}
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate questions: {str(e)}")


@router.post("/job-match")
def analyze_job_match(
    request: ResumeMatchRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Analyze how well a resume matches a job posting."""
    job_description = request.job_description

    if request.job_id:
        # Allow both user-owned jobs and public jobs (user_id is null)
        job = db.query(Job).filter(Job.id == request.job_id).first()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        if job.user_id is not None and job.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this job")
        job_description = job_description or job.job_description

    if not job_description:
        raise HTTPException(status_code=400, detail="Job description is required")
    if not request.resume_text:
        raise HTTPException(status_code=400, detail="Resume text is required")

    try:
        analysis = ai_service.analyze_resume_job_match(
            job_description=job_description,
            resume_text=request.resume_text
        )
        return analysis
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to analyze match: {str(e)}")


@router.post("/improve-resume")
def get_resume_improvements(
    request: ResumeImprovementRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Get suggestions to improve a resume."""
    job_description = request.job_description

    if request.job_id:
        # Allow both user-owned jobs and public jobs (user_id is null)
        job = db.query(Job).filter(Job.id == request.job_id).first()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        if job.user_id is not None and job.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this job")
        job_description = job_description or job.job_description

    if not request.resume_text:
        raise HTTPException(status_code=400, detail="Resume text is required")

    try:
        suggestions = ai_service.get_resume_improvements(
            resume_text=request.resume_text,
            target_job_description=job_description
        )
        return {"suggestions": suggestions}
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate suggestions: {str(e)}")


@router.post("/company-research")
def research_company(
    request: CompanyResearchRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Generate a company research summary for interview prep."""
    company_name = request.company_name
    industry = request.industry
    website = request.website

    if request.company_id:
        # Allow both user-owned companies and public companies (user_id is null)
        company = db.query(Company).filter(Company.id == request.company_id).first()
        if not company:
            raise HTTPException(status_code=404, detail="Company not found")
        if company.user_id is not None and company.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this company")
        company_name = company_name or company.name
        industry = industry or company.industry
        website = website or company.website

    if not company_name:
        raise HTTPException(status_code=400, detail="Company name is required")

    try:
        research = ai_service.generate_company_research(
            company_name=company_name,
            industry=industry,
            website=website
        )
        return {"company_name": company_name, "research": research}
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate research: {str(e)}")


def create_resume_docx(resume_text: str, candidate_name: str = "Resume") -> io.BytesIO:
    """Convert resume text to a formatted DOCX file."""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    lines = resume_text.strip().split('\n')

    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()  # Empty line
            continue

        # Check if it's a section header (all caps or ends with :)
        is_header = (line.isupper() and len(line) < 50) or \
                    (line.endswith(':') and len(line) < 50) or \
                    line.upper() in ['SUMMARY', 'EXPERIENCE', 'SKILLS', 'EDUCATION',
                                     'PROFESSIONAL SUMMARY', 'WORK EXPERIENCE',
                                     'TECHNICAL SKILLS', 'CERTIFICATIONS', 'PROJECTS']

        if is_header:
            p = doc.add_paragraph()
            run = p.add_run(line.upper().rstrip(':'))
            run.bold = True
            run.font.size = Pt(12)
            p.space_after = Pt(6)
        elif line.startswith('•') or line.startswith('-') or line.startswith('*'):
            # Bullet point
            p = doc.add_paragraph(line.lstrip('•-* '), style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25)
        else:
            p = doc.add_paragraph(line)
            # Check if it looks like name (first line, title case, short)
            if lines.index(line) == 0 and len(line) < 50:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(14)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Save to BytesIO
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


@router.post("/ats-resume")
def generate_ats_resume(
    request: ATSResumeRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Generate an ATS-optimized resume tailored to a specific job."""
    job_title = request.job_title
    company_name = request.company_name
    job_description = request.job_description

    if request.job_id:
        # Allow both user-owned jobs and public jobs (user_id is null)
        job = db.query(Job).filter(Job.id == request.job_id).first()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        if job.user_id is not None and job.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this job")
        job_title = job_title or job.title
        job_description = job_description or job.job_description
        if job.company:
            company_name = company_name or job.company.name

    if not job_title:
        raise HTTPException(status_code=400, detail="Job title is required")
    if not job_description:
        raise HTTPException(status_code=400, detail="Job description is required")
    if not request.resume_text:
        raise HTTPException(status_code=400, detail="Resume text is required")

    try:
        ats_resume = ai_service.generate_ats_tailored_resume(
            resume_text=request.resume_text,
            job_title=job_title,
            job_description=job_description,
            company_name=company_name
        )
        return {"ats_resume": ats_resume, "job_title": job_title, "company_name": company_name}
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate ATS resume: {str(e)}")


@router.post("/ats-resume/download")
def generate_ats_resume_docx(
    request: ATSResumeRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Generate an ATS-optimized resume and return as downloadable DOCX file."""
    job_title = request.job_title
    company_name = request.company_name
    job_description = request.job_description

    if request.job_id:
        job = db.query(Job).filter(Job.id == request.job_id).first()
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        if job.user_id is not None and job.user_id != current_user.id:
            raise HTTPException(status_code=403, detail="Not authorized to access this job")
        job_title = job_title or job.title
        job_description = job_description or job.job_description
        if job.company:
            company_name = company_name or job.company.name

    if not job_title:
        raise HTTPException(status_code=400, detail="Job title is required")
    if not job_description:
        raise HTTPException(status_code=400, detail="Job description is required")
    if not request.resume_text:
        raise HTTPException(status_code=400, detail="Resume text is required")

    try:
        # Generate the optimized resume text
        ats_resume = ai_service.generate_ats_tailored_resume(
            resume_text=request.resume_text,
            job_title=job_title,
            job_description=job_description,
            company_name=company_name
        )

        # Convert to DOCX
        docx_file = create_resume_docx(ats_resume)

        # Create filename
        safe_company = re.sub(r'[^\w\s-]', '', company_name or 'Company')[:30]
        safe_title = re.sub(r'[^\w\s-]', '', job_title or 'Resume')[:30]
        filename = f"Resume_{safe_company}_{safe_title}.docx".replace(' ', '_')

        return StreamingResponse(
            docx_file,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate ATS resume: {str(e)}")


# ============== Job-specific endpoints (for job detail page) ==============

from models import UserDocument


def get_user_default_resume(user_id: int, db: Session) -> str:
    """Get the user's default resume text."""
    # First try to find the default resume
    default_resume = db.query(UserDocument).filter(
        UserDocument.user_id == user_id,
        UserDocument.document_type == "resume",
        UserDocument.is_default == True
    ).first()

    if not default_resume:
        # Fall back to most recent resume
        default_resume = db.query(UserDocument).filter(
            UserDocument.user_id == user_id,
            UserDocument.document_type == "resume"
        ).order_by(UserDocument.created_at.desc()).first()

    if not default_resume:
        return None

    return default_resume.content_text


@router.post("/tailor-resume/{job_id}")
def tailor_resume_for_job(
    job_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Generate an ATS-optimized resume tailored for a specific job.
    Uses the user's default resume automatically.
    """
    # Fetch the job
    job = db.query(Job).filter(Job.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    # Allow access to public jobs (user_id is null) or user's own jobs
    if job.user_id is not None and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized to access this job")

    if not job.job_description:
        raise HTTPException(status_code=400, detail="Job description is required for resume tailoring")

    # Get user's default resume
    resume_text = get_user_default_resume(current_user.id, db)
    if not resume_text:
        raise HTTPException(
            status_code=400,
            detail="No resume found. Please upload a resume in the Documents section first."
        )

    # Get company name if available
    company_name = job.company.name if job.company else None

    try:
        # Generate the ATS-optimized resume
        ats_resume = ai_service.generate_ats_tailored_resume(
            resume_text=resume_text,
            job_title=job.title,
            job_description=job.job_description,
            company_name=company_name
        )

        # Convert to DOCX for download
        docx_file = create_resume_docx(ats_resume)

        # Create filename
        safe_company = re.sub(r'[^\w\s-]', '', company_name or 'Company')[:30]
        safe_title = re.sub(r'[^\w\s-]', '', job.title or 'Resume')[:30]
        filename = f"Resume_{safe_company}_{safe_title}.docx".replace(' ', '_')

        return StreamingResponse(
            docx_file,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate tailored resume: {str(e)}")


def create_cover_letter_docx(cover_letter_text: str, job_title: str = "Position", company_name: str = "Company") -> io.BytesIO:
    """Convert cover letter text to a formatted DOCX file."""
    from docx import Document
    from docx.shared import Pt, Inches

    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Add the cover letter content
    paragraphs = cover_letter_text.strip().split('\n\n')

    for para_text in paragraphs:
        if para_text.strip():
            p = doc.add_paragraph(para_text.strip())
            p.paragraph_format.space_after = Pt(12)

    # Save to BytesIO
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream


@router.post("/generate-cover-letter/{job_id}")
def generate_cover_letter_for_job(
    job_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Generate a cover letter tailored for a specific job.
    Uses the user's default resume automatically.
    """
    # Fetch the job
    job = db.query(Job).filter(Job.id == job_id).first()
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")

    # Allow access to public jobs (user_id is null) or user's own jobs
    if job.user_id is not None and job.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Not authorized to access this job")

    if not job.job_description:
        raise HTTPException(status_code=400, detail="Job description is required for cover letter generation")

    # Get user's default resume
    resume_text = get_user_default_resume(current_user.id, db)
    if not resume_text:
        raise HTTPException(
            status_code=400,
            detail="No resume found. Please upload a resume in the Documents section first."
        )

    # Get company name if available
    company_name = job.company.name if job.company else "the company"

    try:
        # Generate the cover letter
        cover_letter = ai_service.generate_cover_letter(
            job_title=job.title,
            company_name=company_name,
            job_description=job.job_description,
            resume_text=resume_text
        )

        # Convert to DOCX for download
        docx_file = create_cover_letter_docx(cover_letter, job.title, company_name)

        # Create filename
        safe_company = re.sub(r'[^\w\s-]', '', company_name or 'Company')[:30]
        safe_title = re.sub(r'[^\w\s-]', '', job.title or 'Position')[:30]
        filename = f"Cover_Letter_{safe_company}_{safe_title}.docx".replace(' ', '_')

        return StreamingResponse(
            docx_file,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except ValueError as e:
        raise HTTPException(status_code=500, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate cover letter: {str(e)}")
